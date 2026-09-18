import os
from pathlib import Path
from dotenv import load_dotenv

load_dotenv()  # must run before scorer / claude_chat read their keys

from flask import Flask, request, jsonify, send_from_directory  # noqa: E402

import scorer  # noqa: E402
import claude_chat  # noqa: E402
from chat_ui import chat_bp  # noqa: E402

ROOT = Path(__file__).parent
PUBLIC = ROOT / "public"
SAMPLE_DIR = ROOT / "sample-data"

app = Flask(__name__, static_folder=str(PUBLIC), static_url_path="")
app.register_blueprint(chat_bp)


@app.route("/")
def index():
    return send_from_directory(PUBLIC, "index.html")


@app.route("/api/health")
def health():
    return jsonify({
        "ok": True,
        "scorer": {"mode": scorer.mode(), "model": scorer._resolved_model or (scorer.MODEL if scorer.gemini_available() else None)},
        "chat": {"mode": claude_chat.mode(), "engine": claude_chat.engine(), "model": claude_chat.engine_model()},
        "gemini_configured": scorer.gemini_available(),
        "anthropic_configured": claude_chat.claude_available(),
        "sample_files": sorted(f.name for f in SAMPLE_DIR.glob("*.md")) if SAMPLE_DIR.exists() else [],
    })


@app.route("/api/sample-data")
def list_samples():
    if not SAMPLE_DIR.exists():
        return jsonify({"files": []})
    return jsonify({"files": sorted(f.name for f in SAMPLE_DIR.glob("*.md"))})


@app.route("/api/sample-data/<filename>")
def get_sample(filename):
    filepath = SAMPLE_DIR / Path(filename).name
    if not filepath.exists() or filepath.suffix != ".md":
        return jsonify({"error": "File not found"}), 404
    return jsonify({"content": filepath.read_text(encoding="utf-8"), "filename": filepath.name})


def _extract_text(filename, blob):
    """Page/slide markers are kept so the scorer can chunk long documents cleanly."""
    import io
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(blob))
        parts = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"[Page {i}]\n{text}")
        return "\n\n".join(parts)

    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(blob))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)
        return "\n".join(parts)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(blob))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        row_text = " | ".join(c.text.strip() for c in row.cells)
                        if row_text.strip(" |"):
                            texts.append(row_text)
            if texts:
                slides.append(f"--- Slide {i} ---\n" + "\n".join(texts))
        return "\n\n".join(slides)

    return blob.decode("utf-8", errors="replace")


@app.route("/api/upload", methods=["POST"])
def upload_file():
    """Accepts one file under 'file', or several under 'files' for batch proposal upload."""
    files = request.files.getlist("files") or ([request.files["file"]] if "file" in request.files else [])
    if not files:
        return jsonify({"error": "No file uploaded"}), 400

    results, errors = [], []
    for f in files:
        try:
            content = _extract_text(f.filename, f.read())
            if not content.strip():
                errors.append({"filename": f.filename, "error": "No readable text found"})
                continue
            results.append({"filename": f.filename, "content": content, "chars": len(content)})
        except Exception as e:
            errors.append({"filename": f.filename, "error": str(e)})

    if not results:
        return jsonify({"error": errors[0]["error"] if errors else "Failed to process file", "errors": errors}), 500
    # Single-file callers keep reading .content / .filename straight off the response.
    return jsonify({**results[0], "files": results, "errors": errors})


@app.route("/api/extract-requirements", methods=["POST"])
def extract_requirements():
    data = request.get_json(silent=True) or {}
    rfp_text = data.get("rfpText", "")
    criteria = data.get("criteria")
    if not rfp_text:
        return jsonify({"error": "RFP text is required."}), 400
    try:
        return jsonify(scorer.extract_requirements(rfp_text, criteria))
    except Exception as e:
        return jsonify({"error": f"Extraction failed: {e}"}), 500


@app.route("/api/analyze", methods=["POST"])
def analyze():
    data = request.get_json(silent=True) or {}
    rfp_text = data.get("rfpText", "")
    proposal_text = data.get("proposalText", "")
    criteria = data.get("criteria") or scorer.DEFAULT_CRITERIA
    requirements = data.get("requirements")
    if not rfp_text or not proposal_text:
        return jsonify({"error": "Both RFP and Proposal text are required."}), 400
    try:
        return jsonify(scorer.analyze(rfp_text, proposal_text, criteria, requirements))
    except scorer.ScorerUnavailable as e:
        return jsonify({"error": str(e), "mode": "mock"}), 503
    except Exception as e:
        print(f"Analysis error: {e}")
        return jsonify({"error": f"Analysis failed: {e}"}), 500


@app.route("/api/suggest-fixes", methods=["POST"])
def suggest_fixes():
    data = request.get_json(silent=True) or {}
    rfp_text = data.get("rfpText", "")
    proposal_text = data.get("proposalText", "")
    analysis = data.get("analysis") or {}
    if not rfp_text or not proposal_text:
        return jsonify({"error": "Both RFP and Proposal text are required."}), 400
    try:
        return jsonify(scorer.suggest_fixes(rfp_text, proposal_text, analysis))
    except scorer.ScorerUnavailable as e:
        return jsonify({"error": str(e), "mode": "mock"}), 503
    except Exception as e:
        return jsonify({"error": f"Suggest-fixes failed: {e}"}), 500


@app.route("/api/validate", methods=["POST"])
def validate():
    """Regression Test Bench — re-runs the known-outcome fixtures and asserts on behaviour."""
    data = request.get_json(silent=True) or {}
    rfp_text = data.get("rfpText", "")
    if not rfp_text:
        return jsonify({"error": "RFP text is required."}), 400

    def load_fixture(name):
        path = SAMPLE_DIR / Path(name).name
        return path.read_text(encoding="utf-8")

    try:
        return jsonify(scorer.run_regression(rfp_text, load_fixture))
    except Exception as e:
        print(f"Regression bench error: {e}")
        return jsonify({"error": f"Regression bench failed: {e}"}), 500


@app.route("/api/bench-fixture/<name>")
def bench_fixture(name):
    """Lets the Review page pull a bench fixture into History for inspection."""
    path = SAMPLE_DIR / Path(name).name
    if not path.exists() or path.suffix != ".md":
        return jsonify({"error": "Fixture not found"}), 404
    return jsonify({"content": path.read_text(encoding="utf-8"), "filename": path.name})


if __name__ == "__main__":
    port = int(os.getenv("PORT", 3000))
    print(f"\n  BidLens running at http://localhost:{port}")
    print(f"  scorer: {scorer.mode()}   assistant: {claude_chat.engine()}")
    if not scorer.gemini_available():
        print("  [i] GEMINI_API_KEY not set - scoring uses the frozen NordFrame Review Mock.")
    if claude_chat.engine() == "gemini":
        print("  [i] ANTHROPIC_API_KEY not set - the assistant runs on the same Gemini key as the scorer.")
    elif claude_chat.engine() == "mock":
        print("  [i] No AI key set - the assistant answers from the frozen review only.")
    print()
    app.run(host="0.0.0.0", port=port, debug=True)
