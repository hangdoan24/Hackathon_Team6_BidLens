import os
import json
import re
from pathlib import Path
from flask import Flask, request, jsonify, send_from_directory
from dotenv import load_dotenv
from google import genai

load_dotenv()

app = Flask(__name__, static_folder="public", static_url_path="")

client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
MODEL = "gemini-3.6-flash"

CHUNK_THRESHOLD = 30000   # ~15 pages — above this, chunk the document
CHUNK_SIZE = 15000        # ~7-8 pages per chunk

SAMPLE_DIR = Path(__file__).parent / "sample-data"
UPLOAD_DIR = Path(__file__).parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

DEFAULT_CRITERIA = [
    {"name": "Problem Understanding", "weight": 15, "description": "Does the proposal correctly reflect the client's actual stated problem/goals from the RFP, not a generic pitch?"},
    {"name": "Scope & Deliverables Clarity", "weight": 15, "description": "Are the deliverables specific and unambiguous? Is it clear what is/isn't included?"},
    {"name": "Pricing Clarity", "weight": 15, "description": "Is pricing clearly stated, broken down, and easy to understand vs. vague or 'on request'?"},
    {"name": "Timeline Clarity", "weight": 10, "description": "Are milestones and dates concrete, not vague ('in due course,' 'as soon as possible')?"},
    {"name": "Completeness vs. RFP Requirements", "weight": 20, "description": "Does the proposal address every requirement the RFP explicitly asked for?"},
    {"name": "Tone & Persuasiveness", "weight": 10, "description": "Does it read as confident, client-focused, and professional - not generic boilerplate?"},
    {"name": "Risk/Assumptions Transparency", "weight": 15, "description": "Are assumptions, dependencies, or risks clearly flagged rather than hidden or omitted?"},
]


# ─── Step 1: Extract structured requirements from RFP ───

def build_extraction_prompt(rfp_text):
    return f"""You are an expert RFP analyst. Extract every explicit requirement from this RFP into a structured list.

For each requirement, identify:
- A short ID (REQ-01, REQ-02, ...)
- The requirement text (what the client is asking for)
- Which section of the RFP it appears in
- Its classification:
  - HARD_CONSTRAINT: explicit, non-negotiable (budget caps, compliance, deadlines, specific technical mandates)
  - MANDATORY: must-have but with some flexibility in how it is met
  - PREFERENCE: nice-to-have, preferred but not required
  - INFORMATIONAL: context or background, not scored

Also identify the client's top priorities — things they emphasize, repeat, or constrain.

Return JSON:
{{
  "requirements": [
    {{
      "id": "REQ-01",
      "requirement": "...",
      "rfp_section": "...",
      "classification": "HARD_CONSTRAINT|MANDATORY|PREFERENCE|INFORMATIONAL"
    }}
  ],
  "client_priorities": [
    {{
      "priority": "...",
      "evidence": "...",
      "rfp_section": "..."
    }}
  ],
  "suggested_weights": [
    {{
      "criterion": "...",
      "weight": <5-30>,
      "reason": "..."
    }}
  ]
}}

RFP:
{rfp_text}"""


# ─── Step 2: Full analysis pipeline ───

def build_analysis_prompt(rfp_text, proposal_text, requirements_json, criteria):
    criteria_list = "\n".join(
        f'{i+1}. **{c["name"]}** (weight: {c["weight"]}): {c["description"]}'
        for i, c in enumerate(criteria)
    )

    req_context = ""
    if requirements_json:
        req_context = f"""
## Pre-extracted RFP Requirements
The following requirements were extracted from the RFP. Use them as your checklist — verify each one against the proposal.
{json.dumps(requirements_json, indent=2)}
"""

    return f"""You are a senior proposal evaluator conducting an evidence-based audit. Your review must be traceable: every claim you make must cite a specific section from the RFP or Proposal.

## Pipeline
Execute these steps in order:

### Step 1 — Requirement Traceability
For each RFP requirement, search the proposal for evidence. Classify each as:
- COVERED: requirement is fully and specifically addressed
- PARTIAL: mentioned but vague, incomplete, or only partly matching
- MISSING: not addressed anywhere in the proposal
- CONTRADICTED: proposal makes a claim that directly conflicts with this requirement
- UNCERTAIN: cannot determine status from available evidence

For each requirement, also provide:
- classification: carry forward from the extracted requirements (HARD_CONSTRAINT, MANDATORY, PREFERENCE, INFORMATIONAL)
- confidence: 0.0-1.0 how confident you are in the status assessment
- blocking: true if this is a HARD_CONSTRAINT that is CONTRADICTED or MISSING

### Step 2 — Risk & Contradiction Detection
Identify red flags:
- Claims that contradict explicit RFP constraints
- Unrealistic timelines or pricing
- Scope that expands beyond what the RFP asked for
- Hidden dependencies or unstated assumptions
- Vague commitments on items the RFP was specific about

Assign severity to each:
- CRITICAL: would likely cause rejection or breach of contract
- HIGH: significant gap that weakens the proposal materially
- MEDIUM: notable issue that should be addressed
- LOW: minor improvement opportunity

### Step 3 — Criteria Scoring (1-5)
Score the proposal on each criterion:
{criteria_list}

Scale:
- 1 = Not addressed or fundamentally wrong
- 2 = Mentioned but vague, incomplete, or generic
- 3 = Partially addressed with some specifics
- 4 = Well addressed with good detail
- 5 = Excellently addressed, specific, and compelling

### Step 4 — Send Readiness
Based on all findings, classify:
- READY: No critical or high issues, coverage > 90%
- REVIEW: Some high issues or coverage 70-90%
- DO_NOT_SEND: Any critical issues, or coverage < 70%

## Evidence Rules
- Every finding MUST include a direct quote or specific section reference from the source document
- If you cannot find evidence for a claim, mark it as UNCERTAIN rather than inventing evidence
- Quote the RFP and Proposal exactly where possible

{req_context}

## Output Format
Return this exact JSON structure:
{{
  "send_readiness": {{
    "status": "READY|REVIEW|DO_NOT_SEND",
    "summary": "<1-2 sentence verdict>",
    "critical_count": <int>,
    "high_count": <int>,
    "medium_count": <int>,
    "coverage_pct": <int 0-100>
  }},
  "requirement_traceability": [
    {{
      "id": "<REQ-XX>",
      "requirement": "<what the RFP asked>",
      "rfp_section": "<section in RFP>",
      "classification": "HARD_CONSTRAINT|MANDATORY|PREFERENCE|INFORMATIONAL",
      "status": "COVERED|PARTIAL|MISSING|CONTRADICTED|UNCERTAIN",
      "severity": "CRITICAL|HIGH|MEDIUM|LOW",
      "confidence": "<0.0-1.0>",
      "blocking": "<true|false>",
      "proposal_evidence": {{
        "section": "<section in proposal or null>",
        "quote": "<exact quote from proposal or null>"
      }},
      "rfp_evidence": {{
        "quote": "<exact quote from RFP>"
      }},
      "business_impact": "<why this matters commercially>",
      "suggested_fix": "<specific, actionable fix>"
    }}
  ],
  "red_flags": [
    {{
      "issue": "<short title>",
      "severity": "CRITICAL|HIGH|MEDIUM|LOW",
      "detail": "<explanation>",
      "rfp_evidence": {{
        "section": "<RFP section>",
        "quote": "<exact quote>"
      }},
      "proposal_evidence": {{
        "section": "<proposal section>",
        "quote": "<exact quote>"
      }},
      "business_impact": "<commercial consequence>",
      "suggested_fix": "<how to resolve>"
    }}
  ],
  "criteria_scores": [
    {{
      "criterion": "<criterion name>",
      "score": <1-5>,
      "comment": "<detailed evidence-based explanation>",
      "strengths": ["<specific strength with evidence>"],
      "weaknesses": ["<specific weakness with evidence>"],
      "citations": [
        {{
          "source": "rfp|proposal",
          "section": "<section name>",
          "quote": "<exact quote>"
        }}
      ],
      "suggested_fixes": ["<actionable fix>"]
    }}
  ],
  "overall": {{
    "score": <1.0-5.0>,
    "verdict": "Strong|Adequate|Needs Work|Weak",
    "summary": "<2-3 sentence executive summary>"
  }}
}}

---

## RFP (Client's Request):
{rfp_text}

---

## PROPOSAL (Draft to Evaluate):
{proposal_text}

---

Analyze now. Return ONLY valid JSON."""


def build_quick_score_prompt(rfp_text, proposal_text):
    return f"""You are a senior proposal evaluator. Quickly score this proposal against the RFP.

Rate each criterion 1-5. Also check if the proposal contradicts any hard constraints (explicit non-negotiable requirements like budget caps, compliance mandates, deadlines).

Rate each criterion 1-5:
1. Problem Understanding — Does the proposal reflect the client's actual stated problem?
2. Scope & Deliverables Clarity — Are deliverables specific and unambiguous?
3. Pricing Clarity — Is pricing clearly stated and broken down?
4. Timeline Clarity — Are milestones and dates concrete?
5. Completeness vs. RFP Requirements — Does it address every RFP requirement?
6. Tone & Persuasiveness — Is it confident, client-focused, professional?
7. Risk/Assumptions Transparency — Are risks and assumptions clearly flagged?

Return ONLY this JSON:
{{{{
  "scores": {{{{
    "Problem Understanding": <1-5>,
    "Scope & Deliverables Clarity": <1-5>,
    "Pricing Clarity": <1-5>,
    "Timeline Clarity": <1-5>,
    "Completeness vs. RFP Requirements": <1-5>,
    "Tone & Persuasiveness": <1-5>,
    "Risk/Assumptions Transparency": <1-5>
  }}}},
  "send_readiness": "READY|REVIEW|DO_NOT_SEND",
  "coverage_pct": <0-100>,
  "critical_count": <int>,
  "high_count": <int>,
  "medium_count": <int>,
  "has_contradictions": <true|false>,
  "has_hard_constraint_contradiction": <true|false>,
  "overall_score": <1.0-5.0>,
  "verdict": "Strong|Adequate|Needs Work|Weak"
}}}}

RFP:
{rfp_text}

Proposal:
{proposal_text}"""


def call_gemini(prompt, max_tokens=8192, temperature=0.3, retries=2):
    last_err = None
    for attempt in range(retries + 1):
        try:
            response = client.models.generate_content(
                model=MODEL,
                contents=prompt,
                config={
                    "response_mime_type": "application/json",
                    "temperature": temperature,
                    "max_output_tokens": max_tokens,
                },
            )
            text = response.text
            try:
                return json.loads(text)
            except json.JSONDecodeError:
                match = re.search(r"\{[\s\S]*\}", text)
                if match:
                    try:
                        return json.loads(match.group())
                    except json.JSONDecodeError:
                        pass
                last_err = ValueError("Failed to parse AI response as JSON")
                if attempt < retries:
                    print(f"JSON parse failed (attempt {attempt + 1}/{retries + 1}), retrying...")
                    continue
                raise last_err
        except Exception as e:
            last_err = e
            if attempt < retries:
                print(f"Gemini call failed (attempt {attempt + 1}/{retries + 1}): {e}, retrying...")
                continue
            raise


# ─── Chunking for large documents ───

def chunk_text(text, chunk_size=CHUNK_SIZE):
    page_pattern = re.compile(r'(?=\[Page \d+\]|--- Slide \d+ ---)')
    segments = page_pattern.split(text)
    segments = [s for s in segments if s.strip()]

    if len(segments) <= 1:
        chunks = []
        remaining = text
        while remaining:
            if len(remaining) <= chunk_size:
                chunks.append(remaining)
                break
            split_at = remaining.rfind('\n\n', 0, chunk_size)
            if split_at == -1 or split_at < chunk_size // 2:
                split_at = remaining.rfind('\n', 0, chunk_size)
            if split_at == -1 or split_at < chunk_size // 2:
                split_at = chunk_size
            chunks.append(remaining[:split_at])
            remaining = remaining[split_at:].lstrip()
        return chunks

    chunks = []
    current = ""
    for seg in segments:
        if len(current) + len(seg) > chunk_size and current:
            chunks.append(current)
            current = seg
        else:
            current += seg
    if current:
        chunks.append(current)
    return chunks


def merge_extraction_results(results):
    all_reqs = []
    all_priorities = []
    weight_sums = {}
    weight_counts = {}

    for result in results:
        all_reqs.extend(result.get("requirements", []))
        all_priorities.extend(result.get("client_priorities", []))
        for sw in result.get("suggested_weights", []):
            crit = sw.get("criterion", "")
            w = sw.get("weight", 10)
            if crit in weight_sums:
                weight_sums[crit] += w
                weight_counts[crit] += 1
            else:
                weight_sums[crit] = w
                weight_counts[crit] = 1

    seen = set()
    unique_reqs = []
    for req in all_reqs:
        key = req.get("requirement", "").strip().lower()[:80]
        if key not in seen:
            seen.add(key)
            unique_reqs.append(req)

    for i, req in enumerate(unique_reqs, 1):
        req["id"] = f"REQ-{i:02d}"

    suggested_weights = []
    for crit, total in weight_sums.items():
        count = weight_counts[crit]
        suggested_weights.append({
            "criterion": crit,
            "weight": round(total / count),
            "reason": f"Averaged from {count} section(s)",
        })

    seen_priorities = set()
    unique_priorities = []
    for p in all_priorities:
        key = p.get("priority", "").strip().lower()[:60]
        if key not in seen_priorities:
            seen_priorities.add(key)
            unique_priorities.append(p)

    return {
        "requirements": unique_reqs,
        "client_priorities": unique_priorities,
        "suggested_weights": suggested_weights,
    }


def extract_requirements_chunked(rfp_text):
    from concurrent.futures import ThreadPoolExecutor, as_completed

    if len(rfp_text) <= CHUNK_THRESHOLD:
        return call_gemini(
            build_extraction_prompt(rfp_text),
            max_tokens=4096,
            temperature=0.2,
        )

    chunks = chunk_text(rfp_text)
    print(f"Large document detected — chunking into {len(chunks)} parts for extraction")

    def extract_chunk(idx, chunk):
        prompt = build_extraction_prompt(chunk).replace(
            "Extract every explicit requirement from this RFP",
            f"Extract every explicit requirement from this section (part {idx + 1}/{len(chunks)}) of a larger RFP",
        )
        return call_gemini(prompt, max_tokens=4096, temperature=0.2)

    results = []
    with ThreadPoolExecutor(max_workers=min(4, len(chunks))) as executor:
        futures = {
            executor.submit(extract_chunk, i, c): i for i, c in enumerate(chunks)
        }
        for future in as_completed(futures):
            try:
                results.append(future.result())
            except Exception as e:
                print(f"Chunk extraction failed: {e}")

    merged = merge_extraction_results(results)
    print(f"Extracted {len(merged['requirements'])} unique requirements from {len(chunks)} chunks")
    return merged


# ─── Routes ───

@app.route("/")
def index():
    return send_from_directory("public", "index.html")


@app.route("/api/sample-data")
def list_samples():
    if not SAMPLE_DIR.exists():
        return jsonify({"files": []})
    files = [f.name for f in SAMPLE_DIR.iterdir() if f.suffix == ".md"]
    return jsonify({"files": sorted(files)})


@app.route("/api/sample-data/<filename>")
def get_sample(filename):
    filepath = SAMPLE_DIR / filename
    if not filepath.exists() or filepath.suffix != ".md":
        return jsonify({"error": "File not found"}), 404
    return jsonify({"content": filepath.read_text(encoding="utf-8")})


@app.route("/api/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    ext = Path(file.filename).suffix.lower()
    try:
        if ext == ".pdf":
            from PyPDF2 import PdfReader
            import io
            reader = PdfReader(io.BytesIO(file.read()))
            parts = []
            for i, page in enumerate(reader.pages, 1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    parts.append(f"[Page {i}]\n{page_text}")
            content = "\n\n".join(parts)
        elif ext == ".docx":
            from docx import Document
            import io
            doc = Document(io.BytesIO(file.read()))
            parts = []
            for para in doc.paragraphs:
                t = para.text.strip()
                if t:
                    parts.append(t)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        parts.append(row_text)
            content = "\n".join(parts)
        elif ext == ".pptx":
            from pptx import Presentation
            import io
            prs = Presentation(io.BytesIO(file.read()))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t:
                                texts.append(t)
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            if row_text.strip(" |"):
                                texts.append(row_text)
                if texts:
                    slides.append(f"--- Slide {i} ---\n" + "\n".join(texts))
            content = "\n\n".join(slides)
        else:
            content = file.read().decode("utf-8")
        return jsonify({"content": content, "filename": file.filename})
    except Exception as e:
        return jsonify({"error": f"Failed to process file: {str(e)}"}), 500


@app.route("/api/extract-requirements", methods=["POST"])
def extract_requirements():
    data = request.get_json()
    rfp_text = data.get("rfpText", "")
    if not rfp_text:
        return jsonify({"error": "RFP text is required."}), 400
    try:
        result = extract_requirements_chunked(rfp_text)
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": f"Extraction failed: {str(e)}"}), 500


@app.route("/api/analyze", methods=["POST"])
def analyze():
    data = request.get_json()
    rfp_text = data.get("rfpText", "")
    proposal_text = data.get("proposalText", "")
    criteria = data.get("criteria", DEFAULT_CRITERIA)
    requirements = data.get("requirements")

    if not rfp_text or not proposal_text:
        return jsonify({"error": "Both RFP and Proposal text are required."}), 400

    try:
        # Step 1: Extract requirements if not provided (chunked for large docs)
        if not requirements:
            try:
                req_result = extract_requirements_chunked(rfp_text)
                requirements = req_result.get("requirements", [])
            except Exception:
                requirements = []

        # Step 2: Full analysis — scale output tokens for large documents
        num_reqs = len(requirements) if requirements else 0
        analysis_tokens = min(16384, max(8192, 8192 + max(0, num_reqs - 15) * 200))
        prompt = build_analysis_prompt(rfp_text, proposal_text, requirements, criteria)
        analysis = call_gemini(prompt, max_tokens=analysis_tokens, temperature=0.3)

        # Weight normalization — auto-normalize to 100%
        total_weight = sum(c["weight"] for c in criteria)
        if total_weight == 0:
            total_weight = 100
        if "criteria_scores" in analysis:
            weighted_sum = 0
            for cs in analysis["criteria_scores"]:
                match_c = next((c for c in criteria if c["name"] == cs.get("criterion")), None)
                w = match_c["weight"] if match_c else total_weight / len(criteria)
                normalized_w = round((w / total_weight) * 100, 1)
                cs["weight"] = normalized_w
                cs["weighted_score"] = round((cs["score"] * w) / total_weight, 2)
                weighted_sum += cs["score"] * w
            analysis["weighted_overall"] = round(weighted_sum / total_weight, 2)

        # Backend-computed Send Readiness with blocker override
        blockers = []
        for req in analysis.get("requirement_traceability", []):
            classification = (req.get("classification") or "").upper()
            status = (req.get("status") or "").upper()
            confidence = req.get("confidence", 0)
            if isinstance(confidence, str):
                try:
                    confidence = float(confidence)
                except ValueError:
                    confidence = 0
            if classification == "HARD_CONSTRAINT" and status == "CONTRADICTED" and confidence >= 0.7:
                blockers.append({
                    "id": req.get("id"),
                    "requirement": req.get("requirement"),
                    "status": status,
                    "confidence": confidence,
                    "rfp_evidence": req.get("rfp_evidence"),
                    "proposal_evidence": req.get("proposal_evidence"),
                })
                req["blocking"] = True

        sr = analysis.get("send_readiness", {})
        base_readiness = sr.get("status", "REVIEW")
        final_readiness = "DO_NOT_SEND" if blockers else base_readiness

        sr["base_readiness"] = base_readiness
        sr["final_readiness"] = final_readiness
        sr["status"] = final_readiness
        sr["blockers"] = blockers
        sr["weighted_score"] = analysis.get("weighted_overall", 0)
        sr["blocker_override"] = bool(blockers)
        if blockers:
            sr["summary"] = f"DO NOT SEND — {len(blockers)} hard constraint(s) contradicted. " + (sr.get("summary") or "")
        analysis["send_readiness"] = sr

        # Attach extracted requirements for frontend
        analysis["extracted_requirements"] = requirements

        return jsonify(analysis)

    except Exception as e:
        print(f"Analysis error: {e}")
        return jsonify({"error": f"Analysis failed: {str(e)}"}), 500


@app.route("/api/validate", methods=["POST"])
def validate():
    data = request.get_json()
    rfp_text = data.get("rfpText", "")
    if not rfp_text:
        return jsonify({"error": "RFP text is required."}), 400

    samples = [
        ("response_1_weak.md", "Weak"),
        ("response_2_medium.md", "Medium"),
        ("response_3_strong.md", "Strong"),
        ("response_4_overpromise.md", "Overpromise"),
    ]

    proposals = {}
    for filename, label in samples:
        path = SAMPLE_DIR / filename
        if path.exists():
            proposals[filename] = {"text": path.read_text(encoding="utf-8"), "label": label}

    from concurrent.futures import ThreadPoolExecutor, as_completed

    def score_one(name, proposal_text):
        prompt = build_quick_score_prompt(rfp_text, proposal_text)
        result = call_gemini(prompt, max_tokens=2048, temperature=0.2)
        return name, result

    results = {}
    try:
        with ThreadPoolExecutor(max_workers=4) as executor:
            futures = {}
            for name, info in proposals.items():
                future = executor.submit(score_one, name, info["text"])
                futures[future] = (name, info["label"])

            for future in as_completed(futures):
                name, label = futures[future]
                try:
                    _, result = future.result()
                    result["label"] = label
                    results[name] = result
                except Exception as e:
                    results[name] = {"label": label, "error": str(e)}
    except Exception as e:
        return jsonify({"error": f"Validation failed: {str(e)}"}), 500

    def get_score(name):
        return results.get(name, {}).get("overall_score", 0)

    weak = get_score("response_1_weak.md")
    medium = get_score("response_2_medium.md")
    strong = get_score("response_3_strong.md")

    overpromise = results.get("response_4_overpromise.md", {})
    checks = [
        {"test": "Strong > Medium", "passed": strong > medium,
         "detail": f"{strong:.1f} vs {medium:.1f}"},
        {"test": "Medium > Weak", "passed": medium > weak,
         "detail": f"{medium:.1f} vs {weak:.1f}"},
        {"test": "Overpromise flagged",
         "passed": overpromise.get("has_contradictions", False),
         "detail": f"contradictions={overpromise.get('has_contradictions', 'N/A')}"},
        {"test": "Overpromise has hard constraint contradiction",
         "passed": overpromise.get("has_hard_constraint_contradiction", False),
         "detail": f"hard_constraint_contradiction={overpromise.get('has_hard_constraint_contradiction', 'N/A')}"},
        {"test": "Weak → DO_NOT_SEND",
         "passed": results.get("response_1_weak.md", {}).get("send_readiness") == "DO_NOT_SEND",
         "detail": results.get("response_1_weak.md", {}).get("send_readiness", "N/A")},
        {"test": "Strong → READY/REVIEW",
         "passed": results.get("response_3_strong.md", {}).get("send_readiness") in ("READY", "REVIEW"),
         "detail": results.get("response_3_strong.md", {}).get("send_readiness", "N/A")},
        {"test": "Overpromise → DO_NOT_SEND",
         "passed": overpromise.get("send_readiness") == "DO_NOT_SEND",
         "detail": overpromise.get("send_readiness", "N/A")},
    ]

    return jsonify({
        "results": results,
        "consistency": {
            "checks": checks,
            "passed": sum(1 for c in checks if c["passed"]),
            "total": len(checks),
        },
    })


if __name__ == "__main__":
    port = int(os.getenv("PORT", 3000))
    print(f"\n  ProposalGuard running at http://localhost:{port}\n")
    app.run(host="0.0.0.0", port=port, debug=True)
